import difflib
import html
import mimetypes
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

import requests


ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path(r"C:\Users\SATORY\Documents")


DOCUMENT_FORMAT_ALIASES = {
    "html": "html",
    "htm": "html",
    "doc": "docx",
    "docx": "docx",
    "dock": "docx",
    "word": "docx",
    "ppt": "pptx",
    "pptx": "pptx",
    "powerpoint": "pptx",
    "presentacion": "pptx",
    "presentación": "pptx",
    "excel": "xlsx",
    "exel": "xlsx",
    "xlsx": "xlsx",
    "xls": "xlsx",
    "txt": "txt",
    "texto": "txt",
}


LANGUAGE_ALIASES = {
    "python": ("python", "py"),
    "py": ("python", "py"),
    "javascript": ("javascript", "js"),
    "js": ("javascript", "js"),
    "typescript": ("typescript", "ts"),
    "ts": ("typescript", "ts"),
    "html": ("html", "html"),
    "css": ("css", "css"),
    "java": ("java", "java"),
    "c#": ("csharp", "cs"),
    "csharp": ("csharp", "cs"),
    "c++": ("cpp", "cpp"),
    "cpp": ("cpp", "cpp"),
    "c": ("c", "c"),
    "php": ("php", "php"),
    "sql": ("sql", "sql"),
    "go": ("go", "go"),
    "golang": ("go", "go"),
    "rust": ("rust", "rs"),
    "ruby": ("ruby", "rb"),
    "kotlin": ("kotlin", "kt"),
    "swift": ("swift", "swift"),
    "r": ("r", "r"),
    "bash": ("bash", "sh"),
    "powershell": ("powershell", "ps1"),
}


@dataclass
class ArtifactRequest:
    kind: str
    formats: list[str] = field(default_factory=list)
    topic: str = ""
    languages: list[str] = field(default_factory=list)
    html_include_assets: bool | None = None
    stage: str = "topic"


def _is_similar_word(word: str, target: str, threshold: float = 0.7) -> bool:
    if word == target:
        return True
    return difflib.SequenceMatcher(None, word, target).ratio() >= threshold


def _contains_similar_word(normalized: str, targets: tuple[str, ...], min_matches: int = 1) -> bool:
    words = re.findall(r"[a-z0-9]+", normalized)
    matches = 0
    for word in words:
        if any(_is_similar_word(word, target) for target in targets):
            matches += 1
            if matches >= min_matches:
                return True
    return False


def _is_web_artifact_request(normalized: str) -> bool:
    web_phrases = (
        "crear pagina web",
        "hacer web",
        "generar sitio",
        "crear sitio web",
        "crear html",
        "pagina web",
        "pajina web",
        "pg web",
        "sitio web",
        "sitio",
        "html",
    )
    creation_terms = ("crear", "hacer", "generar", "diseñar", "construir", "montar", "crear pagina", "haz", "quiero")

    if any(phrase in normalized for phrase in web_phrases) and any(term in normalized for term in creation_terms):
        return True

    if _contains_similar_word(normalized, ("pagina", "pajina", "pag", "pg")) and _contains_similar_word(normalized, ("web", "sitio", "html")):
        return True

    if _contains_similar_word(normalized, ("pagina", "pajina", "pag", "pg")) and any(term in normalized for term in creation_terms):
        # Si el usuario menciona pagina/pajina/pg junto a un pedido de creación, asumir página web si hay contexto web.
        return _contains_similar_word(normalized, ("web", "sitio", "html"))

    return False


def is_ambiguous_artifact_request(text: str) -> bool:
    normalized = _normalize(text)
    ambiguous_phrases = (
        "crear algo",
        "hazme una página",
        "haz un archivo",
        "crear archivo",
        "haz algo",
        "crear algo",
        "quiero algo",
    )
    if any(phrase in normalized for phrase in ambiguous_phrases):
        if _is_web_artifact_request(normalized):
            return False
        if _detect_document_formats(normalized):
            return False
        return True
    return False


def detect_artifact_request(text: str) -> ArtifactRequest | None:
    normalized = _normalize(text)
    # Evitar confundir solicitudes de "tarea"/"tareas"/"pendiente" con creación de documentos
    task_markers = ("tarea", "tareas", "pendiente", "recordatorio", "recordatorios")
    if any(marker in normalized for marker in task_markers):
        return None

    if _is_web_artifact_request(normalized):
        return ArtifactRequest(kind="code", languages=["html", "css", "javascript"], html_include_assets=None)

    code_markers = ("codigo", "programa", "script", "funcion", "clase", "api", "app")
    doc_markers = ("archivo", "documento", "genera", "crear", "crea", "haz", "redacta")

    if any(marker in normalized for marker in code_markers):
        return ArtifactRequest(kind="code", languages=_detect_languages(normalized))

    formats = _detect_document_formats(normalized)
    if any(term in normalized for term in ("documento", "reporte", "ensayo", "pdf", "word")):
        return ArtifactRequest(kind="document", formats=formats or ["docx"])

    if formats or any(marker in normalized for marker in doc_markers):
        return ArtifactRequest(kind="document", formats=formats or ["docx"])

    return None


def parse_languages(text: str) -> list[str]:
    return _detect_languages(_normalize(text)) or ["python"]


def document_prompt(topic: str, formats: list[str]) -> str:
    requested = ", ".join(formats)
    return (
        "Redacta el contenido completo para un documento academico en espanol con formato APA 7.\n"
        f"Tema: {topic}\n"
        f"Formatos de salida que se crearán: {requested}\n"
        "Incluye titulo, resumen, introduccion, desarrollo con subtitulos, conclusion y referencias. "
        "Usa citas parenteticas de ejemplo cuando corresponda. "
        "Entrega solo el contenido final, sin explicar el proceso."
    )


def code_prompt(topic: str, language: str, include_assets: bool = False) -> str:
    if language.lower() == "html" and include_assets:
        return (
            f"Genera un archivo HTML completo para este tema: {topic}.\n"
            "Incluye estilos CSS y JavaScript dentro del mismo archivo HTML usando etiquetas <style> y <script>. "
            "Entrega solo el código HTML final, sin explicaciones adicionales."
        )
    if language.lower() == "css" and include_assets:
        return (
            f"Genera un archivo styles.css para el proyecto web de este tema: {topic}.\n"
            "Incluye un estilo limpio y responsive para una página con encabezados, párrafos y figuras. "
            "Entrega solo el CSS final."
        )
    if language.lower() in {"javascript", "js"} and include_assets:
        return (
            f"Genera un archivo script.js para el proyecto web de este tema: {topic}.\n"
            "Incluye un pequeño comportamiento interactivo como un botón o efectos simples. "
            "Entrega solo el JavaScript final."
        )
    return (
        f"Genera código completo en {language} para este tema: {topic}.\n"
        "Entrega solo el código final, sin explicación externa. "
        "Incluye comentarios breves solo si ayudan a entender el código."
    )


def save_document_artifacts(topic: str, formats: list[str], content: str, html_assets: bool = False) -> list[Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    slug = _slugify(topic)
    paths: list[Path] = []
    clean_formats = _unique(formats or ["docx"])
    images = _download_topic_images(topic, slug)

    for fmt in clean_formats:
        try:
            if fmt == "txt":
                path = OUTPUT_DIR / f"{slug}.txt"
                path.write_text(_to_apa_text(topic, content, images), encoding="utf-8")
            elif fmt == "html":
                if html_assets:
                    paths.extend(_save_document_html_with_assets(topic, content, images))
                    continue
                path = OUTPUT_DIR / f"{slug}.html"
                path.write_text(_to_html(topic, content, images), encoding="utf-8")
            elif fmt == "docx":
                path = OUTPUT_DIR / f"{slug}.docx"
                _save_docx(path, topic, content, images)
            elif fmt == "pptx":
                path = OUTPUT_DIR / f"{slug}.pptx"
                _save_pptx(path, topic, content, images)
            elif fmt == "xlsx":
                path = OUTPUT_DIR / f"{slug}.xlsx"
                _save_xlsx(path, topic, content, images)
            else:
                continue
        except ImportError as exc:
            path = OUTPUT_DIR / f"{slug}_{fmt}_requiere_dependencia.txt"
            path.write_text(
                f"No se pudo crear .{fmt} porque falta una dependencia de Python: {exc}\n"
                "Ejecuta: pip install -r requirements.txt\n",
                encoding="utf-8",
            )
        paths.append(path)

    return paths


def save_code_artifact(topic: str, language: str, content: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    normalized, extension = LANGUAGE_ALIASES.get(_normalize(language), (_slugify(language), "txt"))
    path = OUTPUT_DIR / f"{_slugify(topic)}_{normalized}.{extension}"
    path.write_text(_strip_code_fences(content).strip() + "\n", encoding="utf-8")
    return path


def save_html_bundle(topic: str, html_content: str, css_content: str, js_content: str) -> list[Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    slug = _slugify(topic)
    html_path = OUTPUT_DIR / f"{slug}.html"
    css_path = OUTPUT_DIR / f"{slug}.css"
    js_path = OUTPUT_DIR / f"{slug}.js"
    html_text = _ensure_html_references(html_content, css_path.name, js_path.name)
    html_path.write_text(html_text.strip() + "\n", encoding="utf-8")
    css_path.write_text(_strip_code_fences(css_content).strip() + "\n", encoding="utf-8")
    js_path.write_text(_strip_code_fences(js_content).strip() + "\n", encoding="utf-8")
    return [html_path, css_path, js_path]


def _save_document_html_with_assets(topic: str, content: str, images: list[dict[str, str]]) -> list[Path]:
    slug = _slugify(topic)
    html_path = OUTPUT_DIR / f"{slug}.html"
    css_path = OUTPUT_DIR / f"{slug}.css"
    js_path = OUTPUT_DIR / f"{slug}.js"

    html_text = _to_html_with_external_assets(topic, content, images, css_path.name, js_path.name)
    html_path.write_text(html_text, encoding="utf-8")
    css_path.write_text(_default_html_css(), encoding="utf-8")
    js_path.write_text(_default_html_js(), encoding="utf-8")
    return [html_path, css_path, js_path]


def _ensure_html_references(html_content: str, css_filename: str, js_filename: str) -> str:
    if "<head" in html_content.lower():
        html_text = html_content
    else:
        html_text = f"<!doctype html>\n<html lang=\"es\">\n<head>\n</head>\n<body>\n{html_content}\n</body>\n</html>"
    html_text = html_text.replace("</head>", f"  <link rel=\"stylesheet\" href=\"{css_filename}\">\n  <script src=\"{js_filename}\" defer></script>\n</head>")
    if "</head>" not in html_text:
        html_text = html_text.replace("<html lang=\"es\">", f"<html lang=\"es\">\n<head>\n  <link rel=\"stylesheet\" href=\"{css_filename}\">\n  <script src=\"{js_filename}\" defer></script>\n</head>")
    return html_text


def _to_html_with_external_assets(topic: str, content: str, images: list[dict[str, str]], css_filename: str, js_filename: str) -> str:
    sections: list[str] = []
    headings: list[tuple[str, str]] = []

    for index, image in enumerate(images, start=1):
        image_name = html.escape(Path(image["path"]).name)
        sections.append(
            f"<figure><img src=\"{image_name}\" alt=\"{html.escape(image['title'])}\">"
            f"<figcaption>Figura {index}. {html.escape(image['title'])}. "
            f"Fuente: <a href=\"{html.escape(image['url'])}\">Wikimedia Commons</a>.</figcaption></figure>"
        )

    for i, block in enumerate(_content_blocks(content), start=1):
        if _looks_like_heading(block):
            title = _clean_heading(block)
            hid = f"sec-{i}"
            headings.append((hid, title))
            sections.append(f"<section id=\"{hid}\"><h2>{html.escape(title)}</h2>")
        else:
            if sections and sections[-1].startswith("<section"):
                sections[-1] = sections[-1] + f"<p>{html.escape(block)}</p></section>"
            else:
                sections.append(f"<section><p>{html.escape(block)}</p></section>")

    # Cerrar secciones abiertas que no tengan cierre
    for idx, s in enumerate(sections):
        if s.lstrip().startswith("<section") and not s.rstrip().endswith("</section>"):
            sections[idx] = s + "</section>"

    refs = "".join(f"<li>{html.escape(ref)}</li>" for ref in _apa_image_references(images))

    refs = "".join(f"<li>{html.escape(ref)}</li>" for ref in _apa_image_references(images))

    toc_html = ""
    if headings:
        toc_items = "".join(f"<li><a href=\"#{hid}\">{html.escape(title)}</a></li>" for hid, title in headings)
        toc_html = f"<nav class=\"toc\"><h2>Tabla de contenidos</h2><ul>{toc_items}</ul></nav>"

    return (
        "<!doctype html>\n<html lang=\"es\">\n<head>\n"
        "  <meta charset=\"utf-8\">\n"
        f"  <meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">\n"
        f"  <title>{html.escape(topic.strip().title())}</title>\n"
        f"  <link rel=\"stylesheet\" href=\"{css_filename}\">\n"
        f"  <script src=\"{js_filename}\" defer></script>\n"
        "</head>\n<body>\n"
        f"<div class=\"container\">\n<header><h1>{html.escape(topic.strip().title())}</h1></header>\n"
        + toc_html + "\n<main><article>"
        + "\n".join(sections)
        + "\n</article></main>\n"
        + f"<section><h2>Referencias</h2><ul>{refs}</ul></section>\n"
        + "<footer>Generado por ZENIX · Formato APA 7</footer>\n</div>\n</body>\n</html>\n"
    )


def _default_html_css() -> str:
    return (
        ":root{--bg:#070b14;--panel:#0f172a;--accent:#7c3aed;--muted:#94a3b8;--text:#e6eef8;}\n"
        "body{background:var(--bg);color:var(--text);font-family:Inter,Segoe UI,Helvetica,Arial,serif;margin:0;padding:24px;}\n"
        ".container{max-width:1040px;margin:24px auto;padding:22px;background:linear-gradient(180deg,#071024,#081229);border-radius:12px;box-shadow:0 6px 30px rgba(2,6,23,0.6);}\n"
        "header h1{color:var(--accent);margin:0 0 8px;}\n"
        "nav.toc{margin:18px 0;padding:12px;background:#071026;border-radius:8px;}\n"
        "nav.toc ul{list-style:disc;margin:8px 0 0 18px;padding:0;}\n"
        "article section{margin:18px 0;padding:8px 0;}\n"
        "figure{margin:18px 0;background:#071026;padding:12px;border-radius:10px;}\n"
        "figcaption{color:var(--muted);font-size:0.95rem;}\n"
        "a{color:#60a5fa;text-decoration:none;}\n"
        "a:hover{text-decoration:underline;}\n"
        "footer{border-top:1px solid rgba(255,255,255,0.03);padding-top:12px;margin-top:18px;color:var(--muted);font-size:0.95rem;}\n"
        "@media (max-width:700px){.container{padding:14px;margin:12px}}\n"
    )


def _default_html_js() -> str:
    return (
        "document.addEventListener('DOMContentLoaded', function() {\n"
        "  const btn = document.querySelector('button');\n"
        "  if (btn) {\n"
        "    btn.addEventListener('click', function() {\n"
        "      alert('¡Zenix dice hola! Aquí tienes interacción JS.');\n"
        "    });\n"
        "  }\n"
        "});\n"
    )


def describe_paths(paths: list[Path]) -> str:
    if not paths:
        return "No pude crear archivos con ese formato."
    lines = ["He generado con éxito los siguientes archivos, jefe:"]
    lines.extend(f"- {path}" for path in paths)
    return "\n".join(lines)


def _detect_document_formats(normalized: str) -> list[str]:
    found = []
    for alias, fmt in DOCUMENT_FORMAT_ALIASES.items():
        if re.search(rf"\b{re.escape(alias)}\b", normalized):
            found.append(fmt)
    return _unique(found)


def _detect_languages(normalized: str) -> list[str]:
    found = []
    for alias, (language, _extension) in LANGUAGE_ALIASES.items():
        if re.search(rf"(?<!\w){re.escape(alias)}(?!\w)", normalized):
            found.append(language)
    return _unique(found)


def _download_topic_images(topic: str, slug: str, limit: int = 3) -> list[dict[str, str]]:
    images: list[dict[str, str]] = []
    try:
        response = requests.get(
            "https://commons.wikimedia.org/w/api.php",
            params={
                "action": "query",
                "generator": "search",
                "gsrsearch": topic,
                "gsrnamespace": 6,
                "gsrlimit": limit,
                "prop": "imageinfo",
                "iiprop": "url|extmetadata",
                "format": "json",
            },
            headers={"User-Agent": "ZENIX/2.0 document generator"},
            timeout=25,
        )
        response.raise_for_status()
        pages = response.json().get("query", {}).get("pages", {})
    except Exception:
        return images

    for index, page in enumerate(pages.values(), start=1):
        info = (page.get("imageinfo") or [{}])[0]
        url = info.get("url")
        if not url:
            continue
        try:
            image_response = requests.get(url, timeout=30, headers={"User-Agent": "ZENIX/2.0 document generator"})
            image_response.raise_for_status()
        except Exception:
            continue

        extension = _image_extension(url, image_response.headers.get("content-type", ""))
        if extension not in {".jpg", ".jpeg", ".png"}:
            continue

        image_path = OUTPUT_DIR / f"{slug}_imagen_{index}{extension}"
        image_path.write_bytes(image_response.content)
        metadata = info.get("extmetadata", {})
        title = page.get("title", f"Imagen {index}").replace("File:", "")
        author = _metadata_value(metadata, "Artist") or "Wikimedia Commons"
        license_name = _metadata_value(metadata, "LicenseShortName") or "Licencia Wikimedia Commons"
        images.append({
            "path": str(image_path),
            "title": _clean_html(title),
            "author": _clean_html(author),
            "license": _clean_html(license_name),
            "url": url,
        })
    return images


def _save_docx(path: Path, topic: str, content: str, images: list[dict[str, str]]) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.line_spacing = 2

    document.add_heading(topic.strip().title(), level=0)
    document.add_paragraph("Autor: ZENIX")
    document.add_paragraph("Institucion: Usuario local")
    document.add_paragraph("Formato: APA 7")
    document.add_page_break()

    document.add_heading("Resumen", level=1)
    document.add_paragraph(_make_summary(content))

    if images:
        document.add_heading("Figuras", level=1)
        for index, image in enumerate(images, start=1):
            try:
                document.add_picture(image["path"], width=Inches(5.5))
                document.add_paragraph(f"Figura {index}. {image['title']}. Fuente: {image['url']}")
            except Exception:
                document.add_paragraph(f"Figura {index}. {image['title']}. Fuente: {image['url']}")

    for block in _content_blocks(content):
        if _looks_like_heading(block):
            document.add_heading(_clean_heading(block), level=2)
        else:
            document.add_paragraph(block)
    _add_apa_references_docx(document, images)
    document.save(path)


def _save_pptx(path: Path, topic: str, content: str, images: list[dict[str, str]]) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = topic.strip().title()
    title_slide.placeholders[1].text = "Formato APA 7 | Generado por ZENIX"

    blocks = _content_blocks(content)
    chunks = [blocks[i:i + 4] for i in range(0, len(blocks), 4)] or [["Contenido generado por ZENIX."]]
    for index, chunk in enumerate(chunks[:8], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{topic.strip().title()} {index}"
        body = slide.placeholders[1].text_frame
        body.clear()
        for item in chunk:
            paragraph = body.add_paragraph()
            paragraph.text = _clean_heading(item)
            paragraph.level = 0

    for index, image in enumerate(images[:3], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Figura {index}"
        try:
            slide.shapes.add_picture(image["path"], Inches(1), Inches(1.4), width=Inches(7.5))
        except Exception:
            pass
        textbox = slide.shapes.add_textbox(Inches(1), Inches(6.4), Inches(8), Inches(0.6))
        textbox.text = f"{image['title']}. Fuente: Wikimedia Commons."

    ref_slide = prs.slides.add_slide(prs.slide_layouts[1])
    ref_slide.shapes.title.text = "Referencias"
    ref_slide.placeholders[1].text = "\n".join(_apa_image_references(images)) or "Sin imagenes externas disponibles."
    prs.save(path)


def _save_xlsx(path: Path, topic: str, content: str, images: list[dict[str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Contenido"
    ws["A1"] = "Tema"
    ws["B1"] = topic
    ws["A1"].font = Font(bold=True)
    ws["A3"] = "Seccion"
    ws["B3"] = "Contenido"
    ws["A3"].font = Font(bold=True)
    ws["B3"].font = Font(bold=True)

    for row_index, block in enumerate(_content_blocks(content), start=4):
        ws.cell(row=row_index, column=1).value = row_index - 3
        ws.cell(row=row_index, column=2).value = block
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 90

    ref = wb.create_sheet("Referencias APA")
    ref["A1"] = "Referencias"
    ref["A1"].font = Font(bold=True)
    for row_index, reference in enumerate(_apa_image_references(images), start=2):
        ref.cell(row=row_index, column=1).value = reference
    ref.column_dimensions["A"].width = 120
    wb.save(path)


def _to_html(topic: str, content: str, images: list[dict[str, str]]) -> str:
    # Construir bloques semánticos y lista de encabezados para TOC
    sections: list[str] = []
    headings: list[tuple[str, str]] = []  # (id, title)

    if images:
        for index, image in enumerate(images, start=1):
            image_name = html.escape(Path(image["path"]).name)
            sections.append(
                f"<figure><img src=\"{image_name}\" alt=\"{html.escape(image['title'])}\">"
                f"<figcaption>Figura {index}. {html.escape(image['title'])}. "
                f"Fuente: <a href=\"{html.escape(image['url'])}\">Wikimedia Commons</a>.</figcaption></figure>"
            )

    for i, block in enumerate(_content_blocks(content), start=1):
        if _looks_like_heading(block):
            title = _clean_heading(block)
            hid = f"section-{i}"
            headings.append((hid, title))
            sections.append(f"<section id=\"{hid}\"><h2>{html.escape(title)}</h2>")
        else:
            # añadir párrafo dentro de la última sección si existe, si no crear una
            if sections and sections[-1].startswith("<section"):
                sections[-1] = sections[-1] + f"<p>{html.escape(block)}</p></section>"
            else:
                sections.append(f"<section><p>{html.escape(block)}</p></section>")

    # Cerrar secciones abiertas que no tengan cierre
    for idx, s in enumerate(sections):
        if s.lstrip().startswith("<section") and not s.rstrip().endswith("</section>"):
            sections[idx] = s + "</section>"

    refs = "".join(f"<li>{html.escape(ref)}</li>" for ref in _apa_image_references(images))

    # Generar tabla de contenidos
    toc_html = ""
    if headings:
        toc_items = "".join(f"<li><a href=\"#{hid}\">{html.escape(title)}</a></li>" for hid, title in headings)
        toc_html = f"<nav class=\"toc\"><h2>Tabla de contenidos</h2><ul>{toc_items}</ul></nav>"

    return (
        "<!doctype html>\n<html lang=\"es\">\n<head>\n"
        "  <meta charset=\"utf-8\">\n"
        f"  <meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">\n"
        f"  <title>{html.escape(topic.strip().title())}</title>\n"
        "  <style>"
        ":root{--bg:#070b14;--panel:#0f172a;--accent:#7c3aed;--muted:#94a3b8;--text:#e6eef8}"
        "body{background:var(--bg);color:var(--text);font-family:Inter,Segoe UI,Helvetica,Arial,serif;margin:0;padding:24px;}"
        ".container{max-width:1040px;margin:24px auto;padding:22px;background:linear-gradient(180deg,#071024, #081229);border-radius:12px;box-shadow:0 6px 30px rgba(2,6,23,0.6);}"
        "header h1{color:var(--accent);margin:0 0 8px;}nav.toc{margin:18px 0;padding:12px;background:#071026;border-radius:8px;}nav.toc ul{list-style:disc;margin:8px 0 0 18px;padding:0;}article section{margin:18px 0;padding:8px 0;}figure{margin:18px 0;background:#071026;padding:12px;border-radius:10px;}figcaption{color:var(--muted);font-size:0.95rem;}a{color:#60a5fa;text-decoration:none;}a:hover{text-decoration:underline;}footer{border-top:1px solid rgba(255,255,255,0.03);padding-top:12px;margin-top:18px;color:var(--muted);font-size:0.95rem;}"
        "</style>\n"
        "</head>\n<body>\n"
        f"<div class=\"container\">\n<header><h1>{html.escape(topic.strip().title())}</h1></header>\n"
        + toc_html + "\n<main><article>"
        + "\n".join(sections)
        + "\n</article></main>\n"
        + f"<section><h2>Referencias</h2><ul>{refs}</ul></section>\n"
        + "<footer>Generado por ZENIX · Formato APA 7</footer>\n</div>\n</body>\n</html>\n"
    )



def _to_apa_text(topic: str, content: str, images: list[dict[str, str]]) -> str:
    lines = [
        topic.strip().title(),
        "Autor: ZENIX",
        "Institucion: Usuario local",
        "Formato: APA 7",
        "",
        "Resumen",
        _make_summary(content),
        "",
        content.strip(),
        "",
        "Figuras",
    ]
    if images:
        for index, image in enumerate(images, start=1):
            lines.append(f"Figura {index}. {image['title']}. Fuente: {image['url']}")
    else:
        lines.append("No se pudieron descargar imagenes externas.")
    lines.extend(["", "Referencias"])
    lines.extend(_apa_image_references(images) or ["Sin referencias de imagenes externas."])
    return "\n".join(lines).strip() + "\n"


def _content_blocks(content: str) -> list[str]:
    blocks = [block.strip() for block in re.split(r"\n\s*\n|\r\n\s*\r\n", content.strip()) if block.strip()]
    if len(blocks) <= 1:
        blocks = [line.strip() for line in content.splitlines() if line.strip()]
    return blocks or [content.strip()]


def _looks_like_heading(block: str) -> bool:
    clean = _clean_heading(block)
    return len(clean) <= 80 and (block.startswith("#") or block.endswith(":"))


def _clean_heading(text: str) -> str:
    return text.strip().lstrip("#").strip().rstrip(":").strip()


def _strip_code_fences(content: str) -> str:
    text = content.strip()
    match = re.search(r"```[a-zA-Z0-9_+#-]*\s*(.*?)```", text, flags=re.DOTALL)
    if match:
        return match.group(1).strip()
    return text


def _make_summary(content: str) -> str:
    text = re.sub(r"\s+", " ", content.strip())
    if len(text) <= 600:
        return text
    return text[:600].rsplit(" ", 1)[0].strip() + "."


def _add_apa_references_docx(document, images: list[dict[str, str]]) -> None:
    document.add_heading("Referencias", level=1)
    references = _apa_image_references(images)
    if not references:
        document.add_paragraph("No se pudieron descargar imagenes externas.")
        return
    for reference in references:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.first_line_indent = 0


def _apa_image_references(images: list[dict[str, str]]) -> list[str]:
    references = []
    for image in images:
        references.append(
            f"{image['author']}. (s. f.). {image['title']} [Imagen]. Wikimedia Commons. {image['url']}"
        )
    return references


def _image_extension(url: str, content_type: str) -> str:
    path_extension = Path(urlparse(url).path).suffix.lower()
    if path_extension in {".jpg", ".jpeg", ".png"}:
        return path_extension
    guessed = mimetypes.guess_extension(content_type.split(";")[0].strip())
    return guessed.lower() if guessed else ".jpg"


def _metadata_value(metadata: dict, key: str) -> str:
    value = metadata.get(key, {})
    if isinstance(value, dict):
        return str(value.get("value") or "").strip()
    return str(value or "").strip()


def _clean_html(value: str) -> str:
    text = re.sub(r"<[^>]+>", "", value or "")
    return html.unescape(re.sub(r"\s+", " ", text)).strip()


def _normalize(text: str) -> str:
    replacements = str.maketrans("áéíóúüñ", "aeiouun")
    return text.strip().lower().translate(replacements)


def _slugify(text: str) -> str:
    normalized = _normalize(text)
    slug = re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")
    return slug[:60] or "zenix_archivo"


def _unique(values: list[str]) -> list[str]:
    result = []
    for value in values:
        if value and value not in result:
            result.append(value)
    return result
